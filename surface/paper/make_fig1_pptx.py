#!/usr/bin/env python3
# ==================================================================================================
# make_fig1_pptx.py -- build Figure 1 (fig:overview) of the ICLR 2027 submission as an EDITABLE
# PowerPoint deck.
#
#   python make_fig1_pptx.py                 -> writes figs/fig1.pptx
#   python make_fig1_pptx.py --proxy-png     -> ALSO writes figs/fig1_proxy.png, a matplotlib
#                                               render of the *identical* display list (same
#                                               coordinates, same strings, same colours), used to
#                                               eyeball the layout because this node has no
#                                               PPTX -> PDF renderer.
#
# --------------------------------------------------------------------------------------------
# WHAT THIS IS
# --------------------------------------------------------------------------------------------
# A line-for-line PORT of tex/fig_overview_new.tex -- the approved worked-example mechanism
# figure -- into python-pptx shapes.  The TikZ file is the single source of truth for geometry,
# colour and text; this script reproduces it so the PI can hand-edit the artwork in PowerPoint
# and export a PDF.  Nothing here is an image: every square, chip, slab, icon, badge, arrow and
# label is a real PowerPoint shape, connector or text box.
#
# The figure illustrates ONE toy paired comparison and ONE toy allocation cycle.  Every quantity in
# either is invented for the example; no cell is read from an emitter and no measured number
# appears anywhere:
#
#   (a) eight pairs -> five concordant (grey, no information), three discordant (green) all
#       favouring the candidate -> two pairs curtailed unopened -> verdict: keep
#   (b) bank solve rates 1.00 .80 .60 .50 .75 .00  ->  forecast shades  ->  budget 0, 3, 7, 9, 5, 0
#       ->  groups  ->  three train, one comes back all-equal, two tasks were never funded
#
# 2026-08-29 (morning), PI: "drop panels (b) and (c) -- uninformative; focus on our method; make it
# easy to understand so a reviewer just looks at it and knows our main ideas and contributions."
# DONE.  The old lane (b) (DAPO) and strip (c) (TRACE/VIP) are gone, and with them trVerm, the bin
# icon and the yshift scope.  The height they freed was spent on the drawing rather than pocketed:
# squares, chips, badges and row pitch are all larger.  Four numbered callouts were added: a blue
# disc with a white digit heads each stage column and the two lines beside it are the CONTRIBUTION,
# in five words or fewer, matched to body.tex's \paragraph{Contributions} list -- see the header of
# tex/fig_overview_new.tex, which records the mapping line by line.  They are `disc` ops here: a
# real PowerPoint oval carrying its own text, so the PI can retype them.
#
# 2026-08-29 (03:50), PI, and the reason this file changed today for the second time:
#   * TRIAGE -> BRACE (Bank-calibrated Rollout Allocation under Correlated Episodes).  In the .tex
#     the printed name is the macro \methodname; a .pptx has no macros, so "Brace" is spelled out
#     here in faux small caps, in ONE place (the panel (b) title) -- and nowhere else, because no
#     other string in the artwork names the method.
#   * "DISCORD is our analysis and BRACE is our method built on that analysis."  A SECOND PANEL,
#     (a), now sits above the lane: the paired-outcome worked example the gate of body.tex
#     Sections 2-4 runs on -- concordant pairs grey because they carry no information, discordant
#     pairs green because they are the whole of the evidence, a dashed rule for exact curtailment
#     and two dotted columns curtailment never paid for -- plus one labelled bridge down into the
#     method.  The slide grew from 13.95 x 4.44 cm to 13.95 x 6.00 cm; panel (b)'s coordinates are
#     UNCHANGED, since panel (a) was added above it and TikZ's origin is the bottom-left corner.
#   * The line this file must not blur, and does not: BRACE does not run DISCORD's test.  Nothing
#     is drawn between the green test chip and the blue forecast slab, and the bridge says only
#     that the measurement is the same one level up.
#
# --------------------------------------------------------------------------------------------
# HOW THIS IS ORGANISED (read before tweaking anything)
# --------------------------------------------------------------------------------------------
# * Coordinates ARE the TikZ picture's centimetres: x right, y UP, origin bottom-left, the
#   figure box exactly 13.95 x 6.00 cm (TikZ's \useasboundingbox).  The slide is that size, so
#   the port is 1:1 -- point sizes and line widths below are the PRINTED ones, and cm_to_emu()
#   is the only place the y flip lives.
# * Drawing goes into a backend-neutral display list (`Canvas`), replayed by EITHER the
#   python-pptx backend (real shapes, real text) OR the matplotlib backend (the proxy PNG), so
#   the proxy cannot drift from the .pptx.
# * `tint(colour, pct)` is LaTeX's `colour!pct` (pct% of the colour, rest white) and
#   `mix(a, b, pct)` is `a!pct!b`.  The TikZ file's \colorlet block is transcribed below.
# * Fonts: Arial throughout -- the title 9 pt bold, everything else 7 pt, at PRINT size.  Arial
#   runs about a tenth wider than the TikZ rendition's Times, so the one place where that costs
#   more room than the layout has -- the four callout labels, which are set between fixed discs --
#   is at 6.5 pt (PT_CO), as are panel (a)'s labels, which are set between fixed artwork.  The
#   small-caps names BRACE and DISCORD are faux small caps: leading capital at full size, the
#   rest capitals at 80%.
#
# TO TWEAK: the toy data is in the module-level tables just below; all layout constants are in
# build_figure(), in TikZ centimetres, in the same order as the .tex file.
# ==================================================================================================

from __future__ import annotations

import argparse
import math
import os

# ==================================================================================================
# PALETTE -- tex/head.tex's shared colourblind-safe set, with its fixed meanings.
# ==================================================================================================
# Four of the five: trVerm left with the DAPO lane it belonged to.  Green above, blue below, so
# analysis -> method is a colour change and not only a position.
TR_BLUE  = (0x00, 0x72, 0xB2)   # BRACE / the method
TR_GREEN = (0x00, 0x67, 0x4C)   # DISCORD / the analysis
TR_GREY  = (0x8C, 0x8C, 0x8C)   # NO INFORMATION: a concordant pair, a degenerate group, waste
TR_INK   = (0x26, 0x26, 0x26)   # rules and text
WHITE    = (0xFF, 0xFF, 0xFF)


def mix(a, b, pct):
    """LaTeX's `a!pct!b` -- pct% of a, the rest b."""
    f = pct / 100.0
    return tuple(int(round(x * f + y * (1 - f))) for x, y in zip(a, b))


def tint(c, pct):
    """LaTeX's `c!pct` -- pct% of c mixed into white."""
    return mix(c, WHITE, pct)


# the \colorlet block of fig_overview_new.tex, transcribed
LANE_A = tint(TR_BLUE, 4)
OK_A   = tint(TR_BLUE, 80)      # a solved rollout
ED_A   = tint(TR_BLUE, 85)
DEAD   = tint(TR_GREY, 45)      # an all-equal (degenerate) rollout
ED_D   = tint(TR_GREY, 85)
LANE_D = tint(TR_GREEN, 4)      # panel (a)
OK_D   = tint(TR_GREEN, 80)     # a solved episode in a DISCORDANT pair
ED_G   = tint(TR_GREEN, 85)
TIE    = tint(TR_GREY, 45)      # a solved episode in a CONCORDANT pair -- grey: no information
ED_T   = tint(TR_GREY, 85)
CUT    = tint(TR_GREEN, 70)     # exact curtailment: the rule, and its legend swatch
FLOW_D = tint(TR_GREEN, 62)

C_SUB    = mix(TR_GREY, TR_INK, 75)   # legend and artwork asides (trGrey!75!trInk)
C_CO     = tint(TR_INK, 88)           # a callout's second line   (trInk!88) -- darker than an
                                      # aside, because it is half of a contribution
C_ROWLB  = tint(TR_INK, 80)           # row labels                (trInk!80)
C_FLOW   = tint(TR_INK, 62)           # thin stage arrows         (trInk!62)

# ==================================================================================================
# PAGE, TYPE, UNITS
# ==================================================================================================
FIG_W, FIG_H = 13.95, 6.00      # cm -- exactly TikZ's \useasboundingbox
PT           = 2.54 / 72.27     # one TeX point, in cm (for TikZ's `rounded corners=Npt`)
CM_PT        = 72.0 / 2.54      # one centimetre, in PowerPoint points

FONT     = "Arial"              # available everywhere PowerPoint is; the proxy uses Nimbus Sans,
                                # which is metric-compatible with it.
PT_TITLE = 9.0                  # the title             (TikZ \footnotesize\bfseries)
PT_TEXT  = 7.0                  # everything else       (TikZ \scriptsize)
PT_CO    = 6.5                  # every string the 7 pt Arial rendition outgrows: the four
                                # callout labels, which are set between fixed discs; the legend,
                                # which is laid out from the right edge; and the update node
PT_NUM   = 6.5                  # the white digit inside a callout disc
SC_SMALL = 0.80                 # faux small caps: trailing capitals at 80% of the leading one

LH       = 0.30                 # single-line text-box height, cm
LH_TITLE = 0.32

THIN = (4.0, 3.2)               # TikZ Stealth[length=4pt,width=3.2pt]
BRDG = (4.6, 3.6)               # TikZ Stealth[length=4.6pt,width=3.6pt] -- the (a) -> (b) bridge
FAT  = (5.4, 4.8)               # TikZ Stealth[length=5.4pt,width=4.8pt]

# ==================================================================================================
# THE TOY EXAMPLE.  Kept at module level, in one place, so the four columns cannot drift apart.
#   '#' solved      '.' unsolved      'x' all-equal (grey)
#   '-' this task gets nothing (the short dash)
# ==================================================================================================
# --- panel (a): eight EVALUATED pairs, then two curtailment never paid for.
#     ('incumbent outcome', 'candidate outcome'); '#' solved, '.' unsolved.  A pair whose two
#     outcomes AGREE is concordant and carries nothing; a pair that disagrees is discordant and is
#     the whole of the evidence.  Five concordant, three discordant, all three favouring the
#     candidate -- so three to nil cannot be overturned by the two pairs still unopened, which is
#     exactly the curtailment rule and the reason the last two columns are drawn empty.
PAIR_Y   = (5.37, 5.11)                             # incumbent row, candidate row
PAIRS    = ["##", "..", ".#", "..", "##", ".#", "..", ".#"]
GHOST_X  = [3.46, 3.72]                             # pairs curtailment never evaluated
assert sum(1 for q in PAIRS if q[0] != q[1]) == 3

# --- panel (b) ---------------------------------------------------------------------------------
ROW_Y = [3.10, 2.66, 2.22, 1.78, 1.34, 0.90]        # the six task rows, top to bottom

# (1) the warm bank: banked outcomes of earlier runs -- solve rates 1.00 .80 .60 .50 .75 .00
BANK = ["######", "####.", "##.#.", "#.#.#.", "##.#", "....."]

# (2) the calibrated forecast, P(a group of k comes back all-equal), as a blue tint per task.
#     Monotone in |p - 1/2|: the balanced task is the lightest strip.  None = certified dead.
SHADE = [90, 64, 28, 8, 47, None]

# (3) the allocation that forecast buys: 0 + 3 + 7 + 9 + 5 + 0 = 24 chips
BUDGET = [0, 3, 7, 9, 5, 0]

# (4) the groups that come back, and the verdict on each
GROUPS  = ["-", "##.", "#.##.#.", ".##.#..#.", "xxxxx", "-"]
VERDICT = [None, "ok", "ok", "ok", "x", None]

assert [len(g) for g in GROUPS] == [1 if b == 0 else b for b in BUDGET]
assert len(ROW_Y) == len(BANK) == len(SHADE) == len(BUDGET) == len(GROUPS) == len(VERDICT)


# ==================================================================================================
# BACKEND-NEUTRAL DISPLAY LIST
# ==================================================================================================
# A text "line" is a list of runs.  A run is (text, bold, italic, size_factor, colour_or_None).
def R(text, bold=False, italic=False, sf=1.0, color=None):
    return (text, bold, italic, sf, color)


def SC(word, bold=False, italic=False, color=None):
    """Faux small caps: leading capital at full size, the rest capitalised at SC_SMALL."""
    w = word.strip()
    return [R(w[0].upper(), bold, italic, 1.0, color),
            R(w[1:].upper(), bold, italic, SC_SMALL, color)]


class Canvas:
    """Ordered display list in TikZ centimetres: x right, y UP, origin bottom-left."""

    def __init__(self):
        self.ops = []
        self.dy = 0.0            # TikZ's yshift, applied when an op is RECORDED

    # -- primitives -------------------------------------------------------------------------
    def rect(self, x, y, w, h, fill=None, line=None, lw=0.4, dash=None, radius=0.0):
        """x, y = LOWER-LEFT corner.  radius > 0 -> a rounded rectangle."""
        self.ops.append(dict(kind="rect", x=x, y=y + self.dy, w=w, h=h,
                             fill=fill, line=line, lw=lw, dash=dash, round=radius))

    def ellipse(self, cx, cy, rx, ry, fill=None, line=None, lw=0.4):
        self.ops.append(dict(kind="ellipse", x=cx - rx, y=cy - ry + self.dy, w=2 * rx, h=2 * ry,
                             fill=fill, line=line, lw=lw, dash=None, round=0.0))

    def arc(self, cx, cy, rx, ry, a1, a2, color=TR_INK, lw=0.4):
        """Open elliptical arc, TikZ angles (degrees, counter-clockwise, y up)."""
        self.ops.append(dict(kind="arc", x=cx - rx, y=cy - ry + self.dy, w=2 * rx, h=2 * ry,
                             a1=a1, a2=a2, line=color, lw=lw, fill=None, dash=None, round=0.0))

    def shape(self, name, x, y, w, h, fill=None, line=None, lw=0.4, adj=(), rot=0.0):
        """A PowerPoint preset autoshape: 'can', 'pie', 'homeplate', 'trapezoid'.
        (x, y, w, h) is the UNROTATED bounding box; rot is clockwise degrees about its centre."""
        self.ops.append(dict(kind="shape", name=name, x=x, y=y + self.dy, w=w, h=h,
                             fill=fill, line=line, lw=lw, adj=list(adj), rot=rot))

    def poly(self, pts, fill=None, line=None, lw=0.4, close=True):
        self.ops.append(dict(kind="poly", pts=[(px, py + self.dy) for px, py in pts],
                             fill=fill, line=line, lw=lw, close=close))

    def line(self, x1, y1, x2, y2, color=TR_INK, lw=0.4, dash=None, arrow=None, cap=None):
        self.ops.append(dict(kind="line", x1=x1, y1=y1 + self.dy, x2=x2, y2=y2 + self.dy,
                             color=color, lw=lw, dash=dash, arrow=arrow, cap=cap))

    def curve(self, x1, y1, x2, y2, color=TR_INK, lw=0.4, arrow=None):
        """An S-curve leaving and arriving horizontally -- PowerPoint's curved connector, and
        TikZ's `.. controls ..` fork in lane (b)."""
        self.ops.append(dict(kind="curve", x1=x1, y1=y1 + self.dy, x2=x2, y2=y2 + self.dy,
                             color=color, lw=lw, arrow=arrow))

    # -- a rounded box that carries its own text (ONE editable shape in PowerPoint) ----------
    def node(self, cx, cy, w, h, lines, fill=None, line=None, lw=0.4, dash=None,
             size=PT_TEXT, color=TR_INK, lh=LH, radius=0.0, rot=0, pad_left=0.0):
        self.ops.append(dict(kind="node", x=cx - w / 2.0, y=cy - h / 2.0 + self.dy, w=w, h=h,
                             fill=fill, line=line, lw=lw, dash=dash, round=radius,
                             lines=lines, size=size, color=color, lh=lh, rot=rot,
                             pad_left=pad_left))

    # -- a disc that carries its own text (ONE editable oval in PowerPoint) ------------------
    def disc(self, cx, cy, r, lines, fill=None, line=None, lw=0.4,
             size=PT_NUM, color=WHITE, lh=LH):
        """A numbered-callout marker: TikZ's \\trNum.  Same op shape as `node`, drawn as an oval."""
        self.ops.append(dict(kind="disc", x=cx - r, y=cy - r + self.dy, w=2 * r, h=2 * r,
                             fill=fill, line=line, lw=lw, dash=None, round=0.0,
                             lines=lines, size=size, color=color, lh=lh, rot=0, pad_left=0.0))

    # -- free-standing text -----------------------------------------------------------------
    def text(self, x, y, lines, anchor="nw", align="l", width=6.0,
             size=PT_TEXT, color=TR_INK, lh=LH, fill=None):
        """anchor: which point of the text block sits at (x, y) --
        'nw' top-left, 'n' top-centre, 'w' left-middle, 'e' right-middle, 'c' centre."""
        self.ops.append(dict(kind="text", x=x, y=y + self.dy, lines=lines, anchor=anchor, align=align,
                             width=width, size=size, color=color, lh=lh, fill=fill))

    @staticmethod
    def text_box(op):
        """Resolve a text op to its box (left, bottom, w, h) in figure-cm."""
        h = len(op["lines"]) * op["lh"]
        w = op["width"]
        a = op["anchor"]
        if a == "nw":  left, top = op["x"], op["y"]
        elif a == "n": left, top = op["x"] - w / 2.0, op["y"]
        elif a == "w": left, top = op["x"], op["y"] + h / 2.0
        elif a == "e": left, top = op["x"] - w, op["y"] + h / 2.0
        elif a == "c": left, top = op["x"] - w / 2.0, op["y"] + h / 2.0
        else: raise ValueError(a)
        return left, top - h, w, h


def preset_polygon(name, x, y, w, h, adj, rot):
    """The outline of a polygonal PowerPoint preset, in figure coordinates, after rotation.
    This is the preset's own DrawingML geometry, so the proxy draws the same shape the .pptx
    does rather than an independent approximation."""
    ss = min(w, h)
    if name == "homeplate":                      # <-- the shield: flat back, one point
        dx = adj[0] * ss
        local = [(0, 0), (w - dx, 0), (w, h / 2.0), (w - dx, h), (0, h)]
    else:
        raise ValueError(name)
    pts = [(x + px, y + h - py) for px, py in local]          # preset y is DOWN; figure y is UP
    if not rot:
        return pts
    cx, cy = x + w / 2.0, y + h / 2.0
    t = math.radians(-rot)                                    # pptx rotation is clockwise
    ct, st = math.cos(t), math.sin(t)
    return [(cx + (px - cx) * ct - (py - cy) * st,
             cy + (px - cx) * st + (py - cy) * ct) for px, py in pts]


# ==================================================================================================
# THE SHAPE VOCABULARY of fig_overview_new.tex.  One meaning each, everywhere in the figure.
# ==================================================================================================
def sq(c, x, y, fill, edge):
    """\\trSq -- one rollout outcome (filled = solved, hollow = unsolved)."""
    c.rect(x - 0.11, y - 0.11, 0.22, 0.22,
           fill=fill, line=edge, lw=0.40, radius=0.40 * PT)


def sq_s(c, x, y, fill, edge):
    """\\trSqS -- one EPISODE outcome in a paired comparison, panel (a).  Same rounded square as
    \\trSq, drawn at 0.18 cm instead of 0.22 because it counts a different thing one level down."""
    c.rect(x - 0.09, y - 0.09, 0.18, 0.18,
           fill=fill, line=edge, lw=0.40, radius=0.35 * PT)


def pair(c, x, code):
    """\\trPair -- one evaluated pair, incumbent above candidate.  Grey when the two outcomes
    agree (concordant: no information), green when they disagree (discordant: the evidence)."""
    disc = code[0] != code[1]
    edge = ED_G if disc else ED_T
    solid = OK_D if disc else TIE
    for ch, y in zip(code, PAIR_Y):
        sq_s(c, x, y, solid if ch == "#" else WHITE, edge)


def ghost(c, x):
    """\\trGhost -- a pair exact curtailment never evaluated.  An empty dotted cell, so it reads as
    "never opened" rather than as grey, which in this figure means "opened and uninformative"."""
    for y in PAIR_Y:
        c.rect(x - 0.09, y - 0.09, 0.18, 0.18,
               fill=None, line=tint(TR_GREY, 60), lw=0.35, dash="dot", radius=0.35 * PT)


def chip(c, x, y, fill, edge):
    """\\trChip -- one unit of the generation budget."""
    c.rect(x - 0.08, y - 0.105, 0.16, 0.21,
           fill=fill, line=edge, lw=0.35, radius=0.30 * PT)


def nil(c, x, y, color):
    """\\trNil -- this task gets nothing."""
    c.line(x - 0.09, y, x + 0.09, y, color=color, lw=0.80, cap="round")


def badge_ok(c, x, y, color):
    """\\trBadOK -- verdict badge: this group trains.  Oval + a two-segment check."""
    c.ellipse(x, y, 0.13, 0.13, fill=WHITE, line=color, lw=0.55)
    c.poly([(x - 0.059, y + 0.005), (x - 0.016, y - 0.047), (x + 0.068, y + 0.059)],
           fill=None, line=color, lw=0.70, close=False)


def badge_x(c, x, y, color):
    """\\trBadX -- verdict badge: no gradient.  Oval + a crossed line pair."""
    c.ellipse(x, y, 0.13, 0.13, fill=color, line=color, lw=0.30)
    c.line(x - 0.059, y - 0.059, x + 0.059, y + 0.059, color=WHITE, lw=0.85, cap="round")
    c.line(x - 0.059, y + 0.059, x + 0.059, y - 0.059, color=WHITE, lw=0.85, cap="round")


def num(c, x, y, digit):
    """\\trNum -- a numbered callout marker.  A solid trBlue disc with a white digit, and nothing
    else in the figure is one: the verdict badges are a white disc (trains) or a grey one (dead)."""
    c.disc(x, y, 0.135, [[R(digit, bold=True)]], fill=TR_BLUE, line=TR_BLUE, lw=0.30)


def cylinder(c, x, y, base):
    """\\trCyl -- the warm bank.  MSO_SHAPE.CAN, its lid re-drawn as its own darker oval."""
    c.shape("can", x - 0.125, y - 0.163, 0.250, 0.326,
            fill=tint(base, 12), line=tint(base, 70), lw=0.45, adj=[0.048 / 0.250])
    c.ellipse(x, y + 0.115, 0.125, 0.048, fill=tint(base, 28), line=tint(base, 70), lw=0.45)
    c.arc(x, y + 0.005, 0.125, 0.048, 180, 360, color=tint(base, 45), lw=0.30)


def gauge(c, x, y, base):
    """\\trGauge -- the calibrated forecast.  A half pie, three ticks and a needle."""
    cy = y - 0.055
    c.shape("pie", x - 0.135, cy - 0.135, 0.270, 0.270,
            fill=tint(base, 10), line=tint(base, 70), lw=0.45, adj=[108.0, 216.0])
    for a in (155, 90, 25):
        t = math.radians(a)
        c.line(x + 0.098 * math.cos(t), cy + 0.098 * math.sin(t),
               x + 0.135 * math.cos(t), cy + 0.135 * math.sin(t),
               color=tint(base, 70), lw=0.35)
    t = math.radians(52)
    c.line(x, cy, x + 0.115 * math.cos(t), cy + 0.115 * math.sin(t),
           color=tint(base, 90), lw=0.55, cap="round")
    c.ellipse(x, cy, 0.024, 0.024, fill=tint(base, 90), line=None)


def stack(c, x, y):
    """\\trStack -- unequal allocation: three rectangles of decreasing width."""
    c.rect(x - 0.125, y - 0.115, 0.250, 0.060,
           fill=tint(TR_BLUE, 32), line=tint(TR_BLUE, 70), lw=0.32)
    c.rect(x - 0.125, y - 0.030, 0.170, 0.060,
           fill=tint(TR_BLUE, 20), line=tint(TR_BLUE, 70), lw=0.32)
    c.rect(x - 0.125, y + 0.055, 0.100, 0.060,
           fill=tint(TR_BLUE, 10), line=tint(TR_BLUE, 70), lw=0.32)


def cluster(c, x, y, base):
    """\\trClus -- a group of rollouts, mixed outcomes."""
    r = 0.25 * PT
    c.rect(x - 0.115, y - 0.020, 0.100, 0.100,
           fill=tint(base, 35), line=tint(base, 85), lw=0.30, radius=r)
    c.rect(x + 0.015, y - 0.020, 0.100, 0.100,
           fill=WHITE, line=tint(base, 85), lw=0.30, radius=r)
    c.rect(x - 0.050, y - 0.135, 0.100, 0.100,
           fill=tint(base, 35), line=tint(base, 85), lw=0.30, radius=r)


def shield(c, x, y, base):
    """\\trShield -- certified dead.  A pentagon turned point-down, plus a check.
    (TikZ rounds the point with two Beziers; the preset pentagon is straight-sided.)"""
    c.shape("homeplate", x - 0.1425, y - 0.1075, 0.285, 0.200,
            fill=tint(base, 12), line=tint(base, 85), lw=0.45, adj=[0.122 / 0.200], rot=90)
    c.poly([(x - 0.046, y + 0.010), (x - 0.012, y - 0.030), (x + 0.056, y + 0.050)],
           fill=None, line=tint(base, 85), lw=0.60, close=False)


# -- text shorthands, matching the TikZ node styles -------------------------------------------
def _wide(x, want=3.40):
    """A left-anchored label's box, clamped so it cannot hang off the right edge of the slide.
    Text is left-aligned and never wraps, so the width is cosmetic -- it only decides how big a
    handle the PI grabs in PowerPoint."""
    return max(0.60, min(want, FIG_W - x - 0.05))


def hd(c, x, y, runs):
    """TikZ `hd`: a column heading, anchored west."""
    c.text(x, y, [runs], anchor="w", align="l", width=_wide(x), color=TR_INK)


def sub(c, x, y, txt, lh=LH, size=PT_TEXT):
    """TikZ `sub`: a grey aside (the legend, "no budget", "certified dead"), anchored west."""
    c.text(x, y, [[R(txt)]], anchor="w", align="l", width=_wide(x), color=C_SUB, lh=lh, size=size)


def callout(c, x, y, n, icon, line1, line2):
    """One numbered callout: disc, icon, and the two label lines that name the contribution.
    x is the disc centre; the icon sits 0.32 to its right and the text 0.54, exactly as in
    fig_overview_new.tex, and the label runs on into the air above the next column."""
    num(c, x, y, n)
    icon(c, x + 0.32, y)
    c.text(x + 0.54, y, [[R(line1)]], anchor="w", align="l",
           width=_wide(x + 0.54), size=PT_CO, color=TR_INK)
    c.text(x + 0.54, y - 0.24, [[R(line2)]], anchor="w", align="l",
           width=_wide(x + 0.54), size=PT_CO, color=C_CO)


def rowlabels(c, ys):
    """TikZ `rowlb`: task 1 ... task 6, anchored east at x = 0.72."""
    for i, y in enumerate(ys):
        c.text(0.72, y, [[R("task %d" % (i + 1))]], anchor="e", align="r",
               width=0.72, color=C_ROWLB)


def slab(c, x0, y0, x1, y1, label, fill, line, txt):
    """A tall stage slab with its label turned to read bottom-to-top."""
    c.node((x0 + x1) / 2.0, (y0 + y1) / 2.0, x1 - x0, y1 - y0, [[R(label, bold=True)]],
           fill=fill, line=line, lw=0.40, radius=2 * PT, color=txt, rot=90)


def outcomes(c, code, x0, dx, y, ok, edge):
    """One row of rollout squares from a code string (see the toy-example tables)."""
    if code == "-":
        nil(c, x0 + 0.07, y, ED_D)
        return
    for j, ch in enumerate(code):
        x = x0 + dx * j
        if   ch == "#": sq(c, x, y, ok, edge)
        elif ch == ".": sq(c, x, y, WHITE, edge)
        elif ch == "x": sq(c, x, y, DEAD, ED_D)
        elif ch == "o": sq(c, x, y, WHITE, ED_D)
        else: raise ValueError(ch)


# ==================================================================================================
# THE FIGURE.  Same order, same numbers, as tex/fig_overview_new.tex.
# ==================================================================================================
def build_figure():
    c = Canvas()

    # ====================== panel (a): DISCORD, the measurement ===============================
    # Two rows, one column per evaluated pair.  Everything sits on one horizontal line because the
    # panel is 0.90 cm tall: there is room for artwork OR for a label above it, never both, so the
    # labels are a legend at the right -- GRESO's device, and the same one panel (b) uses.
    c.rect(0, 4.76, FIG_W, 0.90,
           fill=LANE_D, line=tint(TR_GREEN, 45), lw=0.40, radius=3 * PT)
    c.text(0.0, 6.00,
           [[R("(a) ", bold=True)] + SC("Discord", bold=True)
            + [R(": the measurement", bold=True)]],
           anchor="nw", align="l", width=8.0, size=PT_TITLE, lh=LH_TITLE, color=TR_INK)

    for lab, y in zip(("incumbent", "candidate"), PAIR_Y):
        c.text(1.14, y, [[R(lab)]], anchor="e", align="r", width=1.14,
               size=PT_CO, color=C_ROWLB)
    for j, code in enumerate(PAIRS):
        pair(c, 1.34 + 0.26 * j, code)

    # exact curtailment: the rule, then the pairs it never paid for
    c.line(3.32, 4.92, 3.32, 5.56, color=CUT, lw=0.50, dash="dash")
    for gx in GHOST_X:
        ghost(c, gx)

    c.line(3.92, 5.24, 4.16, 5.24, color=FLOW_D, lw=0.55, arrow=THIN)
    c.node(5.15, 5.24, 1.85, 0.56, [[R("exact test on")], [R("discordant pairs")]],
           fill=tint(TR_GREEN, 10), line=tint(TR_GREEN, 75), lw=0.45, radius=2.5 * PT,
           size=PT_CO, lh=0.28, color=TR_INK)
    c.line(6.16, 5.24, 6.44, 5.24, color=FLOW_D, lw=0.55, arrow=THIN)
    badge_ok(c, 6.66, 5.24, ED_G)
    c.text(6.86, 5.24, [[R("keep")]], anchor="w", align="l", width=_wide(6.86), color=TR_INK)

    # panel (a)'s own one-line legend
    sq_s(c, 8.30, 5.36, TIE, ED_T)
    sq_s(c, 8.30, 5.12, TIE, ED_T)
    c.text(8.50, 5.24, [[R("concordant = no signal")]], anchor="w", align="l",
           width=_wide(8.50), size=PT_CO, color=C_SUB)
    c.line(11.15, 5.06, 11.15, 5.42, color=CUT, lw=0.50, dash="dash")
    c.text(11.33, 5.24, [[R("stop early once settled")]], anchor="w", align="l",
           width=_wide(11.33), size=PT_CO, color=C_SUB)

    # ====================== the bridge: the same measurement, one level up =====================
    # ONE line, and it says what is SHARED -- the paired-binary measurement -- not that the method
    # runs the test.  body.tex's abstract draws the line in the same words.
    c.line(3.30, 4.74, 3.30, 4.47, color=tint(TR_BLUE, 80), lw=0.90, arrow=BRDG)
    c.text(3.50, 4.60,
           [[R("same measurement, one level up: an all-equal group carries no gradient")]],
           anchor="w", align="l", width=_wide(3.50, 8.0), size=PT_CO, color=C_CO)

    # ====================== panel (b): BRACE, the method =======================================
    c.rect(0, 0.10, FIG_W, 4.12 - 0.10,
           fill=LANE_A, line=tint(TR_BLUE, 45), lw=0.40, radius=3 * PT)
    c.text(0.0, 4.44,
           [[R("(b) ", bold=True)] + SC("Brace", bold=True)
            + [R(": measure first, then allocate", bold=True)]],
           anchor="nw", align="l", width=8.0, size=PT_TITLE, lh=LH_TITLE, color=TR_INK)

    # one-line legend on the title line.  The three-step ramp is the key to the forecast strips;
    # the three swatches are the rollout squares.  Laid out from the right edge.
    for j, sh in enumerate((14, 50, 88)):
        c.rect(7.83 + 0.24 * j, 4.18, 0.24, 0.22,
               fill=tint(TR_BLUE, sh), line=tint(TR_BLUE, 55), lw=0.35)
    c.text(8.71, 4.29, [[R("P", italic=True), R("(degenerate)")]], anchor="w", align="l",
           width=_wide(8.71), color=C_SUB, lh=0.28, size=PT_CO)
    sq(c, 10.50, 4.29, OK_A, ED_A);   sub(c, 10.66, 4.29, "solved", lh=0.28, size=PT_CO)
    sq(c, 11.57, 4.29, WHITE, ED_A);  sub(c, 11.73, 4.29, "unsolved", lh=0.28, size=PT_CO)
    sq(c, 12.88, 4.29, DEAD, ED_D);   sub(c, 13.04, 4.29, "all-equal", lh=0.28, size=PT_CO)

    # --- the four numbered callouts: one per stage, and they ARE the contribution --------------
    # Each disc sits over its own stage's artwork -- (1) the bank squares, (2) the forecast slab,
    # (3) the budget chips, (4) the groups -- so the discs, not the text, carry the alignment.
    callout(c, 0.98, 3.82, "1", lambda cc, x, y: cylinder(cc, x, y, TR_BLUE),
            "warm bank", "from earlier runs")
    callout(c, 3.48, 3.82, "2", lambda cc, x, y: gauge(cc, x, y, TR_BLUE),
            "calibrated forecast", "correlation priced")
    callout(c, 6.14, 3.82, "3", stack,
            "allocate before generating", "certified-dead excluded")
    callout(c, 9.56, 3.82, "4", lambda cc, x, y: cluster(cc, x, y, TR_BLUE),
            "budget lands where", "gradient exists")

    # --- the three stage slabs every task passes through ---------------------------------------
    for x0, label in ((3.20, "forecast"), (5.30, "allocate"), (8.04, "generate")):
        slab(c, x0, 0.66, x0 + 0.46, 3.34, label,
             fill=tint(TR_BLUE, 20), line=tint(TR_BLUE, 80), txt=tint(TR_BLUE, 85))
    for x1, x2 in ((2.88, 3.16), (4.98, 5.26), (7.72, 8.00)):
        c.line(x1, 2.00, x2, 2.00, color=C_FLOW, lw=0.55, arrow=THIN)

    rowlabels(c, ROW_Y)

    # --- (1) warm bank: unequal amounts of banked evidence -------------------------------------
    for code, y in zip(BANK, ROW_Y):
        outcomes(c, code, 1.20, 0.29, y, OK_A, ED_A)

    # --- (2) the calibrated forecast: one strip per task, shaded by P(all k equal) --------------
    for sh, y in zip(SHADE, ROW_Y):
        if sh is None:
            continue
        c.rect(3.82, y - 0.105, 1.04, 0.21,
               fill=tint(TR_BLUE, sh), line=tint(TR_BLUE, 55), lw=0.40, radius=0.7 * PT)
    c.rect(3.82, 0.795, 1.04, 0.21,
           fill=tint(TR_GREY, 35), line=tint(TR_GREY, 70), lw=0.40, radius=0.7 * PT)
    c.line(3.77, 0.90, 4.91, 0.90, color=tint(TR_GREY, 85), lw=0.70)

    # --- (3) allocate: unequal stacks of chips; nothing for the two dead tasks ------------------
    nil(c, 6.00, 3.10, tint(TR_GREY, 85));  sub(c, 6.18, 3.10, "no budget")
    for b, y in zip(BUDGET, ROW_Y):
        for j in range(b):
            chip(c, 5.96 + 0.195 * j, y, tint(TR_BLUE, 62), ED_A)
    shield(c, 6.04, 0.90, TR_BLUE);         sub(c, 6.24, 0.90, "certified dead")

    # --- (4) generate: the groups that come back, and the verdict on each -----------------------
    for code, y in zip(GROUPS, ROW_Y):
        outcomes(c, code, 8.68, 0.30, y, OK_A, ED_A)
    for v, y in zip(VERDICT, ROW_Y):
        if v == "ok":  badge_ok(c, 11.44, y, tint(TR_BLUE, 85))
        elif v == "x": badge_x(c, 11.44, y, tint(TR_GREY, 85))

    # --- to the optimizer, and the wire that closes the cycle ----------------------------------
    c.line(11.67, 1.34, 11.67, 2.66, color=tint(TR_BLUE, 35), lw=0.60)
    for y in ROW_Y[1:5]:
        c.line(11.57, y, 11.67, y, color=tint(TR_BLUE, 45), lw=0.40)
    c.line(11.67, 2.00, 12.05, 2.00, color=tint(TR_BLUE, 85), lw=1.60, arrow=FAT)
    c.node(12.96, 2.00, 1.75, 0.56, [[R("∇ policy update")]],
           fill=tint(TR_BLUE, 10), line=tint(TR_BLUE, 75), lw=0.45, radius=2.5 * PT, size=PT_CO)

    # the dotted return wire: this cycle's outcomes are the next cycle's evidence
    LOOP = dict(color=tint(TR_BLUE, 65), lw=0.55, dash="dot")
    c.line(12.96, 1.72, 12.96, 0.38, **LOOP)
    c.line(12.96, 0.38, 0.88, 0.38, **LOOP)
    c.line(0.88, 0.38, 0.88, 2.00, **LOOP)
    c.line(0.88, 2.00, 1.07, 2.00, arrow=THIN, **LOOP)
    c.text(6.90, 0.38, [[R("this cycle’s outcomes")]], anchor="c", align="c",
           width=2.45, color=tint(TR_BLUE, 85), fill=LANE_A)

    return c


# ==================================================================================================
# PPTX BACKEND -- real shapes, connectors and text boxes; nothing is an image.
# ==================================================================================================
def render_pptx(c: Canvas, out_path):
    from pptx import Presentation
    from pptx.util import Cm, Pt, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.oxml.ns import qn
    from pptx.oxml import parse_xml

    EMU_CM = 360000
    A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'

    # the ONE place the y flip lives: TikZ cm (y up) -> slide EMU (y down)
    def X(v): return Emu(int(round(v * EMU_CM)))
    def Y(v): return Emu(int(round((FIG_H - v) * EMU_CM)))
    def L(v): return Emu(int(round(v * EMU_CM)))

    PRESET = {"can": MSO_SHAPE.CAN, "pie": MSO_SHAPE.PIE,
              "homeplate": MSO_SHAPE.PENTAGON, "arc": MSO_SHAPE.ARC}

    prs = Presentation()
    prs.slide_width  = Cm(FIG_W)
    prs.slide_height = Cm(FIG_H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])            # blank layout

    DASH = {"dash": MSO_LINE_DASH_STYLE.DASH, "dot": MSO_LINE_DASH_STYLE.ROUND_DOT}

    def no_shadow(shape):
        """PowerPoint's default theme puts a soft shadow on autoshapes; kill it."""
        spPr = shape._element.spPr
        if spPr.find(qn("a:effectLst")) is None:
            spPr.append(parse_xml('<a:effectLst %s/>' % A_NS))

    def style_line(shape, color, lw, dash=None, cap_round=False):
        ln = shape.line
        if color is None:
            ln.fill.background()
            return
        ln.color.rgb = RGBColor(*color)
        ln.width = Pt(lw)
        if dash:
            ln.dash_style = DASH[dash]
        if cap_round:
            ln._get_or_add_ln().set("cap", "rnd")

    def style_fill(shape, color):
        if color is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(*color)

    def add_arrow(shape, spec, lw):
        """python-pptx exposes no arrowhead API; a:tailEnd is last in the <a:ln> sequence.
        OOXML sizes the head as a multiple of the line width, so pick the bucket closest to
        TikZ's absolute Stealth dimensions."""
        def bucket(pts):
            r = pts / max(lw, 0.01)
            return "lg" if r >= 5 else ("med" if r >= 2.5 else "sm")
        shape.line._get_or_add_ln().append(
            parse_xml('<a:tailEnd %s type="stealth" w="%s" len="%s"/>'
                      % (A_NS, bucket(spec[1]), bucket(spec[0]))))

    ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}

    def fill_runs(tf, op_lines, base_size, base_color, lh, align, rot=0, pad_left=0.0):
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = Cm(pad_left)
        tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if rot:
            bp = tf._txBody.find(qn("a:bodyPr"))
            if bp is not None:
                bp.set("vert", "vert270")            # reads bottom-to-top, like TikZ rotate=90
        for i, line in enumerate(op_lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = Pt(lh * CM_PT)
            for (txt, bold, ital, sf, col) in line:
                r = p.add_run()
                r.text = txt
                r.font.name = FONT
                r.font.size = Pt(base_size * sf)
                r.font.bold = bold
                r.font.italic = ital
                r.font.color.rgb = RGBColor(*(col or base_color))

    for op in c.ops:
        k = op["kind"]

        if k in ("rect", "node", "disc", "ellipse", "arc", "shape"):
            if k in ("ellipse", "disc"):
                kind = MSO_SHAPE.OVAL
            elif k == "arc":
                kind = MSO_SHAPE.ARC
            elif k == "shape":
                kind = PRESET[op["name"]]
            elif op["round"] > 0:
                kind = MSO_SHAPE.ROUNDED_RECTANGLE
            else:
                kind = MSO_SHAPE.RECTANGLE
            sh = slide.shapes.add_shape(kind, X(op["x"]), Y(op["y"] + op["h"]),
                                        L(op["w"]), L(op["h"]))
            if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
                sh.adjustments[0] = min(0.5, op["round"] / min(op["w"], op["h"]))
            elif k == "arc":
                # DrawingML angles: clockwise, y DOWN, in 60000ths of a degree; python-pptx
                # normalises adjustments by 100000, so 108.0 == 180 degrees.
                sh.adjustments[0] = ((360 - op["a2"]) % 360) * 0.6
                sh.adjustments[1] = ((360 - op["a1"]) % 360) * 0.6
            elif k == "shape":
                for i, v in enumerate(op["adj"]):
                    sh.adjustments[i] = v
                if op["rot"]:
                    sh.rotation = op["rot"]
            no_shadow(sh)
            style_fill(sh, op["fill"])
            style_line(sh, op["line"], op["lw"], op.get("dash"))
            if k in ("node", "disc"):
                sh.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                fill_runs(sh.text_frame, op["lines"], op["size"], op["color"],
                          op["lh"], PP_ALIGN.CENTER, rot=op["rot"], pad_left=op["pad_left"])

        elif k == "poly":
            pts = op["pts"]
            b = slide.shapes.build_freeform(X(pts[0][0]), Y(pts[0][1]), 1.0)
            b.add_line_segments([(X(px), Y(py)) for px, py in pts[1:]], close=op["close"])
            sh = b.convert_to_shape()
            no_shadow(sh)
            style_fill(sh, op["fill"])
            style_line(sh, op["line"], op["lw"], cap_round=True)

        elif k in ("line", "curve"):
            ctype = MSO_CONNECTOR.STRAIGHT if k == "line" else MSO_CONNECTOR.CURVE
            cn = slide.shapes.add_connector(ctype, X(op["x1"]), Y(op["y1"]),
                                            X(op["x2"]), Y(op["y2"]))
            no_shadow(cn)
            style_line(cn, op["color"], op["lw"], op.get("dash"),
                       cap_round=(op.get("cap") == "round"))
            if op.get("arrow"):
                add_arrow(cn, op["arrow"], op["lw"])

        elif k == "text":
            left, bottom, w, h = Canvas.text_box(op)
            tb = slide.shapes.add_textbox(X(left), Y(bottom + h), L(w), L(h))
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            if op["fill"] is not None:
                style_fill(tb, op["fill"])
            fill_runs(tb.text_frame, op["lines"], op["size"], op["color"],
                      op["lh"], ALIGN[op["align"]])

        else:
            raise ValueError(k)

    prs.save(out_path)
    return out_path


# ==================================================================================================
# MATPLOTLIB PROXY -- the same display list, at the same 1:1 print size.  For eyeballing only.
# ==================================================================================================
def render_proxy(c: Canvas, out_path, dpi=300):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import (FancyBboxPatch, Rectangle, Ellipse, Polygon,
                                    FancyArrowPatch, Wedge, Arc)
    from matplotlib.path import Path

    PROXY_FONT = "Nimbus Sans"      # metric-compatible with Arial (Helvetica widths)

    def rgb(t): return tuple(v / 255.0 for v in t)

    fig = plt.figure(figsize=(FIG_W / 2.54, FIG_H / 2.54), dpi=dpi)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, FIG_W); ax.set_ylim(0, FIG_H)
    ax.set_axis_off(); ax.set_facecolor("white")
    DASHES = {"dash": (2.4, 1.4), "dot": (0.7, 1.3)}

    fig.canvas.draw()
    rend = fig.canvas.get_renderer()

    def run_width(txt, bold, ital, size):
        t = ax.text(0, 0, txt, fontname=PROXY_FONT, fontsize=size,
                    fontweight="bold" if bold else "normal",
                    fontstyle="italic" if ital else "normal")
        bb = t.get_window_extent(renderer=rend); t.remove()
        return bb.width / dpi * 2.54

    def draw_lines(lines, left, top, w, size, color, lh, align, rot=0, cx=None, cy=None):
        for i, line in enumerate(lines):
            widths = [run_width(t, b, it, size * sf) for (t, b, it, sf, _) in line]
            total = sum(widths)
            if rot:                       # bottom-to-top: lay the runs out along +y
                y = cy - total / 2.0
                for (t, b, it, sf, col), wd in zip(line, widths):
                    ax.text(cx, y, t, fontname=PROXY_FONT, fontsize=size * sf,
                            fontweight="bold" if b else "normal",
                            fontstyle="italic" if it else "normal",
                            color=rgb(col or color), ha="left", va="center",
                            rotation=90, rotation_mode="anchor", zorder=6)
                    y += wd
                continue
            x = left + (w - total) / 2.0 if align == "c" else (
                left + w - total if align == "r" else left)
            ycen = top - (i + 0.5) * lh
            for (t, b, it, sf, col), wd in zip(line, widths):
                ax.text(x, ycen - lh * 0.07, t, fontname=PROXY_FONT, fontsize=size * sf,
                        fontweight="bold" if b else "normal",
                        fontstyle="italic" if it else "normal",
                        color=rgb(col or color), ha="left", va="center", zorder=6)
                x += wd

    def s_path(x1, y1, x2, y2):
        """PowerPoint's curved connector: an S leaving and arriving horizontally."""
        xm = (x1 + x2) / 2.0
        return Path([(x1, y1), (xm, y1), (xm, y2), (x2, y2)],
                    [Path.MOVETO, Path.CURVE4, Path.CURVE4, Path.CURVE4])

    for op in c.ops:
        k = op["kind"]

        if k in ("rect", "node", "disc", "ellipse", "poly", "shape", "arc"):
            fc = rgb(op["fill"]) if op.get("fill") else "none"
            ec = rgb(op["line"]) if op.get("line") else "none"
            kw = dict(linewidth=op["lw"], zorder=2, fc=fc, ec=ec)
            if op.get("dash"):
                kw["linestyle"] = (0, DASHES[op["dash"]])
            if k in ("ellipse", "disc"):
                p = Ellipse((op["x"] + op["w"] / 2, op["y"] + op["h"] / 2),
                            op["w"], op["h"], **kw)
            elif k == "arc":
                kw.pop("fc")
                ax.add_patch(Arc((op["x"] + op["w"] / 2, op["y"] + op["h"] / 2),
                                 op["w"], op["h"], theta1=op["a1"], theta2=op["a2"],
                                 zorder=3, edgecolor=ec, linewidth=op["lw"]))
                continue
            elif k == "shape":
                if op["name"] == "pie":                       # the gauge's half disc
                    a1 = (360 - op["adj"][1] / 0.6) % 360
                    a2 = (360 - op["adj"][0] / 0.6) % 360
                    p = Wedge((op["x"] + op["w"] / 2, op["y"] + op["h"] / 2),
                              op["w"] / 2, a1, a2, **kw)
                elif op["name"] == "can":                     # the warm bank cylinder
                    rx, ry = op["w"] / 2.0, op["adj"][0] * min(op["w"], op["h"])
                    cx0 = op["x"] + rx
                    ybot, ytop = op["y"] + ry, op["y"] + op["h"] - ry
                    ax.add_patch(Ellipse((cx0, ybot), 2 * rx, 2 * ry, fc=fc, ec="none", zorder=2))
                    ax.add_patch(Rectangle((op["x"], ybot), op["w"], ytop - ybot,
                                           fc=fc, ec="none", zorder=2))
                    ax.add_patch(Arc((cx0, ybot), 2 * rx, 2 * ry, theta1=180, theta2=360,
                                     edgecolor=ec, linewidth=op["lw"], zorder=3))
                    for xw in (op["x"], op["x"] + op["w"]):
                        ax.plot([xw, xw], [ybot, ytop], color=ec, lw=op["lw"], zorder=3)
                    ax.add_patch(Ellipse((cx0, ytop), 2 * rx, 2 * ry,
                                         fc=fc, ec=ec, lw=op["lw"], zorder=3))
                    continue
                else:
                    p = Polygon(preset_polygon(op["name"], op["x"], op["y"], op["w"], op["h"],
                                               op["adj"], op["rot"]), closed=True, **kw)
            elif k == "poly":
                if op["close"]:
                    p = Polygon(op["pts"], closed=True, **kw)
                else:
                    xs = [q[0] for q in op["pts"]]; ys = [q[1] for q in op["pts"]]
                    ax.plot(xs, ys, color=ec, lw=op["lw"], solid_capstyle="round",
                            solid_joinstyle="round", zorder=5)
                    continue
            elif op["round"] > 0:
                r = op["round"]
                p = FancyBboxPatch((op["x"] + r, op["y"] + r), op["w"] - 2 * r, op["h"] - 2 * r,
                                   boxstyle="round,pad=%f" % r, **kw)
            else:
                p = Rectangle((op["x"], op["y"]), op["w"], op["h"], **kw)
            ax.add_patch(p)
            if k in ("node", "disc"):
                draw_lines(op["lines"], op["x"] + op["pad_left"],
                           op["y"] + op["h"] / 2 + len(op["lines"]) * op["lh"] / 2,
                           op["w"] - op["pad_left"], op["size"], op["color"], op["lh"], "c",
                           rot=op["rot"], cx=op["x"] + op["w"] / 2, cy=op["y"] + op["h"] / 2)

        elif k in ("line", "curve"):
            col = rgb(op["color"]); lw = op["lw"]
            ls = (0, DASHES[op["dash"]]) if op.get("dash") else "-"
            if op.get("arrow"):
                style = "-|>,head_length=%.2f,head_width=%.2f" % (op["arrow"][0],
                                                                  op["arrow"][1])
                kw = dict(arrowstyle=style, mutation_scale=1.0, color=col, lw=lw,
                          linestyle=ls, shrinkA=0, shrinkB=0, zorder=3)
                if k == "curve":
                    ax.add_patch(FancyArrowPatch(path=s_path(op["x1"], op["y1"],
                                                             op["x2"], op["y2"]), **kw))
                else:
                    ax.add_patch(FancyArrowPatch((op["x1"], op["y1"]),
                                                 (op["x2"], op["y2"]), **kw))
            elif k == "curve":
                from matplotlib.patches import PathPatch
                ax.add_patch(PathPatch(s_path(op["x1"], op["y1"], op["x2"], op["y2"]),
                                       fc="none", ec=col, lw=lw, zorder=3))
            else:
                ax.plot([op["x1"], op["x2"]], [op["y1"], op["y2"]], color=col, lw=lw,
                        linestyle=ls, zorder=3,
                        solid_capstyle="round" if op.get("cap") == "round" else "butt")

        elif k == "text":
            left, bottom, w, h = Canvas.text_box(op)
            if op["fill"] is not None:
                ax.add_patch(Rectangle((left, bottom), w, h, fc=rgb(op["fill"]),
                                       ec="none", zorder=4))
            draw_lines(op["lines"], left, bottom + h, w,
                       op["size"], op["color"], op["lh"], op["align"])

    fig.savefig(out_path, dpi=dpi, facecolor="white")
    plt.close(fig)
    return out_path


# ==================================================================================================
def main():
    here = os.path.dirname(os.path.abspath(__file__))
    ap = argparse.ArgumentParser(description="Build Figure 1 as an editable .pptx.")
    ap.add_argument("--out", default=os.path.join(here, "figs", "fig1.pptx"))
    ap.add_argument("--proxy-png", action="store_true",
                    help="also write figs/fig1_proxy.png from the identical display list")
    ap.add_argument("--dpi", type=int, default=300, help="proxy PNG resolution")
    args = ap.parse_args()

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    canvas = build_figure()
    render_pptx(canvas, args.out)
    print("wrote", args.out, "(%d shapes)" % len(canvas.ops))
    if args.proxy_png:
        p = os.path.splitext(args.out)[0] + "_proxy.png"
        render_proxy(canvas, p, dpi=args.dpi)
        print("wrote", p, "(proxy render, matplotlib -- NOT the pptx)")


if __name__ == "__main__":
    main()
